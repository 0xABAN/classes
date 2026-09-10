# /// script
# requires-python = ">=3.12"
# dependencies = ["python-pptx>=1.0,<2", "numpy>=2,<3", "pillow>=11,<13"]
# ///
"""Build a repeated-sampling lesson with `uv run build_deck.py`.

Text, data points, and curves remain editable in PowerPoint. The only raster
asset is an original decorative gradient. Checks run whenever the deck builds.
"""

from io import BytesIO
from pathlib import Path

import numpy as np
from numpy.polynomial import Polynomial
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
PAPER = "F5F5F4"
INK = "181D16"
MUTED = "686C66"
RULE = "D4D6D0"
ORANGE = "C65326"
PURPLE = "7963AF"
FAINT_PURPLE = "C4B8DC"
FONT = "Arial"
NOISE_STD = .4
Y_LIMITS = (-.65, 2.45)


def text(slide, value, x, y, w, h, size=24, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line_value in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line_value
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(0)
    return box


def line(slide, x1, y1, x2, y2, color=RULE, width=1, dashed=False):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width)
    if dashed:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shape


def path(slide, points, color, width=2.5, dashed=False):
    """Keep each plotted curve a single editable PowerPoint vector shape."""
    coords = [(round(x * 1000), round(y * 1000)) for x, y in points]
    builder = slide.shapes.build_freeform(*coords[0], scale=Inches(1) / 1000)
    builder.add_line_segments(coords[1:], close=False)
    shape = builder.convert_to_shape()
    shape.fill.background()
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width)
    if dashed:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shape


def dot(slide, x, y, color=INK, hollow=False, r=.045):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x-r), Inches(y-r), Inches(r*2), Inches(r*2)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(PAPER if hollow else color)
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(1.5)


def gradient():
    """An original soft color field; no source-site artwork or font assets."""
    stops = np.array([[196, 175, 240], [244, 167, 149], [247, 183, 108], [235, 207, 128]])
    yy = np.linspace(0, 3, 900)
    rgb = np.stack([np.interp(yy, np.arange(4), stops[:, c]) for c in range(3)], axis=1)
    pixels = np.repeat(rgb[:, None, :], 320, axis=1)
    pixels += np.random.default_rng(9).normal(0, 1.4, pixels.shape[:2])[:, :, None]
    image = Image.fromarray(np.clip(pixels, 0, 255).astype(np.uint8))
    stream = BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream


def base(prs, n, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(PAPER)
    slide.shapes.add_picture(gradient(), Inches(12.87), Inches(.22), Inches(.24), Inches(7.06))
    text(slide, title, .62, 1.05, 12, .8, 38)
    text(slide, f"{n:02d}/07", 11.94, 7.12, .75, .22, 11, MUTED)
    return slide


def project(rect, a, b):
    """All regression plots use the same x and y limits."""
    x, y, w, h = rect
    return x + (a+1)/2*w, y+h-(b-Y_LIMITS[0])/(Y_LIMITS[1]-Y_LIMITS[0])*h


def plot(slide, rect, curves=(), dots=()):
    """Split out-of-range curves rather than flattening or rescaling them."""
    x, y, w, h = rect
    line(slide, x, y+h, x+w, y+h)
    line(slide, x, y, x, y+h)
    text(slide, "x", x+w-.1, y+h+.10, .7, .3, 14, MUTED)
    text(slide, "y", x-.30, y-.2, 1, .3, 14, MUTED)
    for xx, yy, color, dashed, width in curves:
        segment = []
        for a, b in zip(xx, yy):
            if Y_LIMITS[0] <= b <= Y_LIMITS[1]:
                segment.append(project(rect, a, b))
            else:
                if len(segment) > 1:
                    path(slide, segment, color, width, dashed)
                segment = []
        if len(segment) > 1:
            path(slide, segment, color, width, dashed)
    for xx, yy, color, hollow in dots:
        for a, b in zip(xx, yy):
            if -1 <= a <= 1 and Y_LIMITS[0] <= b <= Y_LIMITS[1]:
                dot(slide, *project(rect, a, b), color, hollow)


def legend(slide, items, x=.92, y=6.45):
    for label, color, kind in items:
        if kind in ("dot", "hollow"):
            dot(slide, x+.07, y+.1, color, kind == "hollow")
        else:
            line(slide, x, y+.1, x+.28, y+.1, color, 2, kind == "dash")
        text(slide, label, x+.39, y-.03, 2.7, .35, 13, MUTED)
        x += 3.04


def note(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def truth(x):
    return .3 + .4*x + 1.3*x*x


def data():
    """Refit every degree on the same eight samples; sample A keeps the old seed."""
    x = np.linspace(-1, 1, 20)
    samples = np.array([
        truth(x) + np.random.default_rng(seed).normal(0, NOISE_STD, len(x))
        for seed in (12, 20, 21, 22, 23, 24, 25, 26)
    ])
    models = {
        degree: [Polynomial.fit(x, sample, degree) for sample in samples]
        for degree in (1, 2, 12)
    }
    return x, samples, models


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = "The bias–variance tradeoff"
    prs.core_properties.subject = "CMPSC 448 — bias and variance through repeated training"
    prs.core_properties.author = "CMPSC 448"
    x, samples, models = data()
    g = np.linspace(-1, 1, 400)
    predictions = {degree: np.array([model(g) for model in fits]) for degree, fits in models.items()}
    means = {degree: values.mean(axis=0) for degree, values in predictions.items()}
    two = [(left, 3.1, 5.0, 2.85) for left in (.98, 7.12)]
    large = (.98, 2.94, 7.12, 3.10)

    s = base(prs, 1, "New data, same underlying pattern")
    for rect, sample, label in zip(two, samples[:2], ("Sample A", "Sample B")):
        text(s, label, rect[0], 2.65, 4.9, .3, 16)
        plot(s, rect, dots=[(x, sample, INK, False)])
    legend(s, [("training observations", INK, "dot")])
    note(s, f"Open with a question: If we collected our training data again, would we get exactly the same observations? Both panels use the same 20 x positions and the same underlying relationship, y = 0.3 + 0.4x + 1.3x² + noise. Only the independent Gaussian output noise changes, with standard deviation {NOISE_STD}. Sample A is the original fixed-seed dataset; sample B is another draw. This is a controlled thought experiment, not train versus test and not distribution shift. We have not fitted a model yet. The other presentation explains underfitting and overfitting; our focus is what changes across training samples and why it matters for prediction.")

    s = base(prs, 2, "Different training data changes the model")
    fixed_x = .5
    fixed_predictions = [float(model(fixed_x)) for model in models[12][:2]]
    for rect, sample, model, label, prediction in zip(
        two, samples[:2], models[12][:2], ("Sample A", "Sample B"), fixed_predictions
    ):
        text(s, f"{label}  /  prediction {prediction:.2f}", rect[0], 2.65, 4.9, .3, 16)
        plot(s, rect, [(g, model(g), ORANGE, False, 3)], [(x, sample, INK, False)])
        marker_x, marker_y = project(rect, fixed_x, prediction)
        line(s, marker_x, rect[1], marker_x, rect[1]+rect[3], PURPLE, 1, True)
        dot(s, marker_x, marker_y, ORANGE, False, .075)
        text(s, "x = 0.5", marker_x-.38, 6.05, 1, .3, 14, PURPLE)
    legend(s, [("training observations", INK, "dot"), ("degree-12 fit", ORANGE, "line")], y=6.62)
    note(s, f"Use the exact same samples A and B from slide 1, and fit the same degree-12 polynomial procedure separately to each. Keep the inputs, model family, and fitting method fixed. At x = {fixed_x}, A predicts {fixed_predictions[0]:.2f} and B predicts {fixed_predictions[1]:.2f}. The vertical guides mark that identical input. The orange dots are model predictions, not additional observations. Retraining changed the prediction even though the underlying problem did not change. Two samples illustrate the idea; they do not estimate population variance reliably. We now repeat the experiment eight times.")

    s = base(prs, 3, "High variance: predictions change a lot")
    for rect, degree, label in zip(two, (2, 12), ("Degree 2  /  less variation", "Degree 12  /  more variation")):
        text(s, label, rect[0], 2.65, 5, .3, 16)
        curves = [(g, values, FAINT_PURPLE, False, 1.5) for values in predictions[degree]]
        plot(s, rect, curves + [(g, means[degree], ORANGE, False, 2.8)])
    legend(s, [("one fit per sample", PURPLE, "line"), ("average of 8 fits", ORANGE, "line")])
    note(s, "Each pale curve is a model fitted to one of eight training samples, including A and B. Every degree sees exactly the same eight samples. The orange curve is the pointwise arithmetic average of those eight predictions, not a ninth fit and not the true relationship. At any fixed input, variance describes the spread of predictions around their average across training samples. The degree-12 procedure is more sensitive here; the degree-2 procedure is more stable. Wiggliness by itself is not the definition of variance, and variance is not the scatter of the raw data. Eight curves are a teaching illustration, not a precise population estimate. Axis limits are identical; curve portions outside the visible range are omitted. Flexibility need not produce high variance in every setting; sample size and regularization also matter.")

    s = base(prs, 4, "High bias: predictions are consistently off")
    curves = [(g, values, FAINT_PURPLE, False, 1.5) for values in predictions[1]]
    plot(s, large, curves + [
        (g, truth(g), INK, True, 2),
        (g, means[1], ORANGE, False, 3),
    ])
    # A gap at a fixed input distinguishes bias from the spread of the pale fits.
    mean_at_zero = float(np.mean([model(0) for model in models[1]]))
    gap_x, truth_y = project(large, 0, truth(0))
    _, mean_y = project(large, 0, mean_at_zero)
    line(s, gap_x, truth_y, gap_x, mean_y, ORANGE, 2)
    dot(s, gap_x, mean_y, ORANGE, False, .065)
    dot(s, gap_x, truth_y, INK, True, .065)
    text(s, "Bias", 8.75, 3.35, 3.6, .6, 30, ORANGE)
    text(s, "Average prediction\nminus truth.", 8.75, 4.12, 3.7, 1.15, 23)
    legend(s, [("one straight fit", PURPLE, "line"), ("average of 8 fits", ORANGE, "line"), ("true pattern", INK, "dash")])
    note(s, f"Now reveal the true regression function, shown as the dashed curve. We know it because this is synthetic data; it is usually unknown in practice. These are degree-1 fits to the same eight samples. They cluster, yet their average misses the curved relationship. At x = 0 the average shown predicts {mean_at_zero:.2f}, while the true mean output is {truth(0):.2f}; the vertical segment marks that gap. Bias at a fixed input is the expected model prediction over training samples minus the true regression function at that input. The displayed eight-fit average is only a finite-sample approximation to that expectation. Bias is not a single model's residual against one noisy observation. Low variance does not guarantee low bias. A straight-line family is too restrictive for this quadratic relationship, connecting briefly to the other student's underfitting example.")

    s = base(prs, 5, "A good model stays close to the truth")
    curves = [(g, values, FAINT_PURPLE, False, 1.5) for values in predictions[2]]
    plot(s, large, curves + [
        (g, truth(g), INK, True, 2),
        (g, means[2], ORANGE, False, 3),
    ])
    text(s, "Low bias.\nLow variance.", 8.75, 3.57, 3.8, 1.4, 28, ORANGE)
    legend(s, [("one degree-2 fit", PURPLE, "line"), ("average of 8 fits", ORANGE, "line"), ("true pattern", INK, "dash")])
    note(s, "Return to the degree-2 fits from slide 3, now with the true pattern revealed. All data, noise levels, and axes remain unchanged. Their average is close to the true relationship, unlike the degree-1 average; their spread is smaller than the degree-12 spread. Low bias and low variance are relative descriptions for this controlled example, not claims of zero prediction error. Because the quadratic family contains the true function and noise has mean zero, fixed-design least squares is unbiased here in expectation. The eight-fit average need not equal truth exactly. Increasing flexibility from a line to a quadratic reduces systematic error; increasing it further can add sensitivity without useful signal. That is the tradeoff, not a universal requirement that bias and variance must always move in opposite directions.")

    s = base(prs, 6, "Even good models make errors")
    text(s, "Expected squared error", .95, 2.8, 5.55, .8, 31)
    text(s, "=", 6.25, 2.8, .65, .8, 31, MUTED)
    text(s, "bias²", 7.0, 2.8, 1.25, .8, 31, ORANGE)
    text(s, "+", 8.25, 2.8, .5, .8, 31, MUTED)
    text(s, "variance", 8.8, 2.8, 2.0, .8, 31, PURPLE)
    text(s, "+", 10.8, 2.8, .5, .8, 31, MUTED)
    text(s, "noise", 11.3, 2.8, 1.25, .8, 31)
    for label, description, left, color in [
        ("BIAS²", "Systematic error.", .98, ORANGE),
        ("VARIANCE", "Sensitivity to the sample.", 5.02, PURPLE),
        ("NOISE", "Uncertainty that remains.", 9.06, INK),
    ]:
        text(s, label, left, 4.38, 3.5, .3, 13, color)
        text(s, description, left, 4.9, 3.5, 1.0, 23)
    text(s, "Even the true pattern cannot predict every noisy observation.", .98, 6.18, 11.6, .55, 22, MUTED)
    note(s, f"This equality is the squared-error decomposition at a fixed input x: E[(f_hat_D(x) - Y_new)²] = (E_D[f_hat_D(x)] - f(x))² + Var_D[f_hat_D(x)] + Var(Y_new | x). Average over random training datasets D and independent fresh output noise, with Y_new = f(x) + epsilon and E[epsilon | x] = 0. The noise term is conditional variance, not standard deviation; here it is {NOISE_STD}² = {NOISE_STD**2:.2f}. Even a predictor equal to the true mean function cannot predict the random noise in a fresh observation. The decomposition is not a formula for one observed residual or one training loss, and is not generally valid in this form for classification error or every loss function. Noise is irreducible relative to the information available to the model; better measurements or extra informative features can change that information. The build checks the finite-ensemble squared-error identity numerically. Do not imply the displayed eight curves are the full population expectation.")

    s = base(prs, 7, "We want accurate, stable predictions")
    text(s, "BIAS", 1.0, 3.2, 4.8, .35, 14, ORANGE)
    text(s, "Is the average\nprediction off?", 1.0, 3.9, 5.0, 1.6, 32)
    text(s, "VARIANCE", 7.12, 3.2, 4.8, .35, 14, PURPLE)
    text(s, "How much does\nretraining change it?", 7.12, 3.9, 5.0, 1.6, 32)
    note(s, "Ask the audience to distinguish two cases: all retrained models agree but miss the true mean, versus retrained models disagree strongly at the same input. The first signals systematic bias; the second signals variance. A single fitted curve cannot directly show either population quantity. Underfitting and overfitting are useful outcomes discussed by the other presenter; bias and variance explain different sources of prediction error across repeated training. Close with the two questions on screen. Teaching reference: An Introduction to Statistical Learning, 2nd edition, section 2.2.2, The Bias-Variance Trade-Off, https://www.statlearning.com/. All data and plots are original synthetic examples; the visual language is inspired by the saved superlogical profile without source-site branding or assets.")

    # Native shapes inherit a shadow from the template unless disabled explicitly.
    for slide in prs.slides:
        for effect in slide._element.xpath(".//a:effectRef"):
            effect.set("idx", "0")

    verify(prs, x, samples, models, g, predictions, means)
    output = ROOT / "bias-variance-tradeoff.pptx"
    prs.save(output)
    saved = Presentation(output)
    assert len(saved.slides) == 7
    assert all(slide.has_notes_slide and slide.notes_slide.notes_text_frame.text for slide in saved.slides)
    assert not any(shape.has_table for slide in saved.slides for shape in slide.shapes)
    # Compare actual saved dot positions so the refit cannot silently change data.
    assert np.all((samples[:2] >= Y_LIMITS[0]) & (samples[:2] <= Y_LIMITS[1]))
    sample_positions = []
    for slide in list(saved.slides)[:2]:
        dots = [shape for shape in slide.shapes if shape.name.startswith("Oval")
                and shape.fill.fore_color.rgb == RGBColor.from_string(INK)][:40]
        assert len(dots) == 40
        sample_positions.append([(shape.left, shape.top) for shape in dots])
    assert sample_positions[0] == sample_positions[1]
    assert abs(fixed_predictions[0] - fixed_predictions[1]) > .4
    for number, slide in enumerate(saved.slides, 1):
        visible = [shape.text for shape in slide.shapes if shape.has_text_frame]
        assert f"{number:02d}/07" in visible
        assert any(shape.left == Inches(12.87) and shape.width == Inches(.24) for shape in slide.shapes)
        assert not any(shape.top == Inches(7.02) for shape in slide.shapes)
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= saved.slide_width + Inches(.02)
            assert shape.top + shape.height <= saved.slide_height + Inches(.02)
    print(f"Saved and checked: {output}")
    print("Fixed-input predictions at x = 0.5:", fixed_predictions)


def verify(prs, x, samples, models, g, predictions, means):
    """Check the actual repeated-sampling claims and the squared-error identity."""
    assert len(prs.slides) == 7
    assert samples.shape == (8, 20)
    assert not np.array_equal(samples[0], samples[1])
    for degree, fits in models.items():
        for sample, fit in zip(samples, fits):
            np.testing.assert_allclose(fit(x), Polynomial.fit(x, sample, degree)(x))
    variances = {degree: values.var(axis=0) for degree, values in predictions.items()}
    squared_bias = {degree: (mean-truth(g))**2 for degree, mean in means.items()}
    assert variances[12].mean() > variances[2].mean()
    assert variances[12].mean() > variances[1].mean()
    assert squared_bias[1].mean() > squared_bias[2].mean()
    for degree, values in predictions.items():
        expected_error = ((values-truth(g))**2).mean(axis=0) + NOISE_STD**2
        decomposition = squared_bias[degree] + variances[degree] + NOISE_STD**2
        np.testing.assert_allclose(expected_error, decomposition, rtol=1e-12, atol=1e-12)
    print("Mean empirical variance:", {degree: float(value.mean()) for degree, value in variances.items()})
    print("Mean squared empirical bias:", {degree: float(value.mean()) for degree, value in squared_bias.items()})


if __name__ == "__main__":
    build()
