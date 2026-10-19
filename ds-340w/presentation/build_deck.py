"""Rebuild six editable slides using the supplied PowerPoint's actual package.

Keep its size, masters, theme, background, and decorative image on each slide.
Replace the lesson-specific shapes and notes. Rendered examples come from the
project's drawing interpreter, not from a separate illustration generator.
"""

from io import BytesIO
import json
from pathlib import Path
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent))
from pixelcode.dataset import build_records
from pixelcode.renderer import render

INK, MUTED, ORANGE, PURPLE, RULE = "181D16", "686C66", "C65326", "7963AF", "D4D6D0"


def text(slide, value, x, y, width, height, size=24, color=INK, bold=False, link=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.line_spacing = 1.2
        paragraph.space_after = Pt(0)
        if link:
            for run in paragraph.runs:
                run.hyperlink.address = link
                run.font.color.rgb = RGBColor.from_string(color)
                run.font.underline = False
    return shape


def rule(slide, x, y, width):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                       Inches(x + width), Inches(y))
    shape.line.color.rgb = RGBColor.from_string(RULE)
    shape.line.width = Pt(1)
    shape._element.spPr.append(OxmlElement("a:effectLst"))


def sprite(slide, code, x, y, size=2.3):
    # Crop to the depicted object for legibility; the dataset retains the full canvas.
    image = render(code)
    pixels = [(xx, yy) for yy in range(32) for xx in range(32)
              if image.getpixel((xx, yy)) != (245, 245, 244)]
    bounds = (min(p[0] for p in pixels), min(p[1] for p in pixels),
              max(p[0] for p in pixels) + 1, max(p[1] for p in pixels) + 1)
    image = image.crop(bounds).resize((320, 320), Image.Resampling.NEAREST)
    stream = BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    slide.shapes.add_picture(stream, Inches(x), Inches(y), Inches(size), Inches(size))


def build():
    content = json.loads((ROOT / "content.json").read_text())
    notes = json.loads((ROOT / "notes.json").read_text())
    assert len(content["slides"]) == len(notes) == 6
    deck = Presentation(ROOT / "template.pptx")
    while len(deck.slides) > 6:
        item = deck.slides._sldIdLst[-1]
        deck.part.drop_rel(item.rId)
        deck.slides._sldIdLst.remove(item)

    for number, (slide, item, note) in enumerate(zip(deck.slides, content["slides"], notes), 1):
        for shape in list(slide.shapes):
            # Every template slide has one narrow gradient picture on the right.
            if shape.shape_type != 13:
                shape._element.getparent().remove(shape._element)
        text(slide, item["title"], .62, 1.05, 12, .8, 36)
        text(slide, f"{number:02d}/06", 11.94, 7.12, .75, .22, 11, MUTED)
        seconds = content["timing_seconds"][number - 1]
        slide.notes_slide.notes_text_frame.text = f"Target: {seconds} seconds.\n\n{note}"

    rows = build_records()
    example = next(r for r in rows if r["object"] == "mushroom" and r["split"] == "train")
    original = example["reference_code"]
    edited = original.replace(example["color_hex"], "#b8579e")
    blue = original.replace(example["color_hex"], "#437ab5")

    slide, item = deck.slides[0], content["slides"][0]
    text(slide, item["subtitle"], .66, 1.95, 11, .4, 20, MUTED)
    text(slide, item["goal"], .98, 2.65, 11, .5, 28)
    sprite(slide, original, 1.35, 3.5, 2.1)
    sprite(slide, edited, 7.15, 3.5, 2.1)
    text(slide, "→", 4.65, 4.0, 1.3, .8, 42, ORANGE)
    text(slide, item["before"], 1.35, 5.82, 3.5, .4, 20)
    text(slide, item["after"], 6.5, 5.82, 5.8, .45, 20)
    text(slide, item["caption"], .98, 6.45, 11, .3, 13, MUTED)
    text(slide, item["outline"], .98, 6.9, 10.7, .3, 12, MUTED)

    slide, item = deck.slides[1], content["slides"][1]
    for i, paper in enumerate(item["papers"]):
        y = 2.25 + i * 1.25
        text(slide, paper["name"], .98, y, 3, .45, 25, PURPLE, link=paper["url"])
        text(slide, paper["venue"], .98, y + .48, 3, .3, 16, MUTED)
        text(slide, paper["idea"], 4.35, y, 7.7, .45, 25)
        text(slide, paper["transfer"], 4.35, y + .50, 7.7, .4, 20, MUTED)
        if i < 2:
            rule(slide, .98, y + 1.03, 11.1)
    text(slide, item["takeaway"], .98, 6.38, 11.3, .55, 19)

    slide, item = deck.slides[2], content["slides"][2]
    text(slide, item["question"], .98, 2.02, 11.2, .95, 27)
    text(slide, item["prompt"], .98, 3.12, 11, .4, 20, MUTED)
    sprite(slide, blue, 1.65, 3.8, 1.65)
    sprite(slide, original, 7.65, 3.8, 1.65)
    text(slide, item["left_label"], .98, 5.65, 5.5, .5, 20)
    text(slide, item["right_label"], 6.9, 5.65, 5.3, .5, 20)
    text(slide, item["hypothesis"], .98, 6.32, 11.2, .6, 18)
    text(slide, item["caption"], .98, 7.02, 10.7, .24, 11, MUTED)

    slide, item = deck.slides[3], content["slides"][3]
    text(slide, item["model"], .98, 2.12, 11, .5, 24, PURPLE)
    for i, (step, detail) in enumerate(zip(item["steps"], item["details"])):
        x = .98 + 2.98*i
        text(slide, f"{i+1:02d}", x, 3.15, 1, .5, 24, ORANGE)
        text(slide, step, x, 3.91, 2.65, .85, 25)
        text(slide, detail, x, 4.94, 2.6, .85, 18, MUTED)
        if i < 3:
            text(slide, "→", x + 2.45, 3.22, .4, .5, 22, MUTED)
    text(slide, item["takeaway"], .98, 6.13, 11.2, .5, 22)
    text(slide, item["status"], .98, 6.83, 10.8, .35, 13, MUTED)

    slide, item = deck.slides[4], content["slides"][4]
    text(slide, item["count"], .98, 2.35, 5, .65, 37, PURPLE)
    text(slide, item["split"], .98, 3.17, 5, .5, 25)
    text(slide, item["description"], .98, 3.94, 5.5, .5, 20)
    text(slide, item["canvas"], .98, 4.55, 5.5, .5, 18, MUTED)
    text(slide, item["baseline"], 6.9, 2.35, 5.2, .85, 25)
    text(slide, item["comparison"], 6.9, 3.45, 5.2, .8, 21, PURPLE)
    for i, label in enumerate(item["metrics"]):
        text(slide, label, 6.9, 4.48 + i*.65, 5.2, .5, 18)
    text(slide, item["warning"], .98, 6.52, 11.2, .55, 20)

    slide, item = deck.slides[5], content["slides"][5]
    for i, target in enumerate(item["targets"]):
        text(slide, target, .98 + i*6, 2.4, 5.6, .65, 32, ORANGE if i == 0 else PURPLE)
    text(slide, item["caveat"], .98, 3.35, 11.2, .45, 19, MUTED)
    text(slide, item["ready"], .98, 4.25, 11, .5, 25)
    text(slide, item["next"], .98, 5.02, 11, .5, 25)
    text(slide, item["repo_label"], .98, 6.22, 3, .4, 20, PURPLE, link=item["repo_url"])
    text(slide, item["data_label"], 5.2, 6.22, 3, .4, 20, PURPLE, link=item["data_url"])
    text(slide, "github.com/0xABAN/classes/tree/main/ds-340w", .98, 6.80, 10.7, .32, 14, MUTED,
         link=item["repo_url"])

    deck.core_properties.title = "Training AI to Pixelate with Code"
    deck.core_properties.subject = "DS 340W midterm proposal"
    deck.core_properties.author = "Adam Torres Encarnacion"
    deck.core_properties.keywords = "DS 340W, Team 7, midterm, pixel art, preference learning"
    deck.save(ROOT / "training-ai-to-pixelate.pptx")
    transcript = []
    for i, (item, note, duration) in enumerate(zip(content["slides"], notes, content["timing_seconds"]), 1):
        transcript.append(f"Slide {i}: {item['title']} ({duration} seconds)\n\n{note}")
    (ROOT / "speaker-notes.txt").write_text("\n\n".join(transcript) + "\n")
    print("Built six slides; timing totals 5 minutes background + 5 minutes method/evaluation.")


if __name__ == "__main__":
    build()
